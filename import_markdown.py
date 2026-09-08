#!/usr/bin/env python3
# RESANTA CRM v23.6.87 · import markdown/discounted goods from 1C email
import os, re, json, io, imaplib, email, hashlib
from datetime import datetime, timezone, timedelta, date
from email.header import decode_header, make_header
from email.utils import parsedate_to_datetime
from zoneinfo import ZoneInfo

import requests
from openpyxl import load_workbook

IMAP_HOST=os.getenv("IMAP_HOST","imap.gmail.com")
IMAP_PORT=int(os.getenv("IMAP_PORT","993"))
IMAP_USER=os.environ["IMAP_USER"].strip()
# Google app passwords are often copied with spaces / NBSP. Existing CRM importers normalize them too.
IMAP_PASS="".join(os.environ["IMAP_PASS"].split()).replace("\xa0","")
SUPABASE_URL=os.environ["SUPABASE_URL"].strip().rstrip("/")
SUPABASE_KEY=os.environ["SUPABASE_KEY"].strip()
LOOKBACK_DAYS=int(os.getenv("MARKDOWN_LOOKBACK_DAYS","14"))
SUBJECT_KEYS=[x.strip().lower() for x in os.getenv("MARKDOWN_SUBJECT_KEYS","уценк,crm").split(",") if x.strip()]
FULL_SNAPSHOT=os.getenv("MARKDOWN_FULL_SNAPSHOT","true").lower() in ("1","true","yes","y")
MINSK=ZoneInfo("Europe/Minsk")

ALIASES={
 "source_key":["id","ид","ключ","уникальный id","уникальный идентификатор","идентификатор","uid"],
 "article":["артикул","код","код товара","sku","номенклатурный номер"],
 "nomenclature":["номенклатура","товар","наименование","наименование товара"],
 "quantity":["количество","кол-во","остаток","остаток уценки","qty"],
 "base_price":["цена","базовая цена","цена с ндс","розничная цена","цена продажи","price"],
 "source_type":["причина","тип возврата","источник","откуда","вид уценки"],
 "source_document":["документ","документ возврата","номер документа","документ сервиса"],
 "source_date":["дата","дата возврата","дата документа"],
 "serial_no":["серийный номер","серийник","serial","s/n","sn"],
 "source_comment":["комментарий","что с товаром","дефект","состояние","описание дефекта","примечание"],
 "warranty_active":["гарантия","гарантия действует","с гарантией"],
 "warranty_months":["гарантия мес","гарантия месяцев","гарантия, мес","срок гарантии"],
 "warranty_note":["условия гарантии","комментарий гарантии"]
}

def log(s): print(s,flush=True)

def norm(v):
    s=str(v or "").strip().lower().replace("ё","е")
    s=re.sub(r"\s+"," ",s)
    return s

def decoded(v):
    try:return str(make_header(decode_header(v or "")))
    except:return str(v or "")

def num(v,default=0):
    if v is None or v=="": return default
    if isinstance(v,(int,float)): return float(v)
    s=str(v).strip().replace("\xa0","").replace(" ","").replace(",",".")
    s=re.sub(r"[^0-9.\-]","",s)
    try:return float(s)
    except:return default

def int_or_none(v):
    n=num(v,None)
    return None if n is None else int(round(n))

def bool_warranty(v):
    if v is None or str(v).strip()=="": return True
    s=norm(v)
    if s in ("нет","не","0","false","без гарантии","не действует"): return False
    return True

def iso_date(v,fallback):
    if isinstance(v,datetime): return v.date().isoformat()
    if isinstance(v,date): return v.isoformat()
    s=str(v or "").strip()
    if not s:return fallback.isoformat()
    for fmt in ("%d.%m.%Y","%Y-%m-%d","%d/%m/%Y"):
        try:return datetime.strptime(s[:10],fmt).date().isoformat()
        except:pass
    return fallback.isoformat()

def choose_header(ws):
    best=None
    alias_terms={norm(x) for vals in ALIASES.values() for x in vals}
    for r in range(1,min(ws.max_row,20)+1):
        vals=[norm(ws.cell(r,c).value) for c in range(1,min(ws.max_column,50)+1)]
        score=sum(1 for x in vals if x in alias_terms)
        if best is None or score>best[0]:best=(score,r,vals)
    if not best or best[0]<2:
        raise RuntimeError("Не нашёл строку заголовков. Нужны минимум Артикул/Номенклатура и Цена/Количество.")
    return best[1]

def colmap(headers):
    m={}
    for key,aliases in ALIASES.items():
        aset={norm(x) for x in aliases}
        for idx,h in enumerate(headers):
            if norm(h) in aset:
                m[key]=idx
                break
    if "article" not in m or "nomenclature" not in m:
        raise RuntimeError("В файле обязательны колонки «Артикул» и «Номенклатура».")
    return m

def cell(row,m,key,default=None):
    i=m.get(key)
    return row[i] if i is not None and i<len(row) else default

def parse_xlsx(content,report_date,filename):
    wb=load_workbook(io.BytesIO(content),data_only=True,read_only=True)
    ws=wb[wb.sheetnames[0]]
    hr=choose_header(ws)
    headers=[ws.cell(hr,c).value for c in range(1,ws.max_column+1)]
    m=colmap(headers)
    rows=[]
    for rnum,row in enumerate(ws.iter_rows(min_row=hr+1,values_only=True),start=hr+1):
        article=str(cell(row,m,"article","") or "").strip()
        name=str(cell(row,m,"nomenclature","") or "").strip()
        if not article and not name: continue
        if not article or not name:
            log(f"⚠️ Строка {rnum} пропущена: нет артикула или номенклатуры")
            continue
        doc=str(cell(row,m,"source_document","") or "").strip()
        serial=str(cell(row,m,"serial_no","") or "").strip()
        comment=str(cell(row,m,"source_comment","") or "").strip()
        explicit=str(cell(row,m,"source_key","") or "").strip()
        if explicit:
            skey=explicit
        else:
            stable="|".join([article,serial,doc,comment])
            if stable.replace("|",""):
                skey="AUTO-"+hashlib.sha1(stable.encode("utf-8")).hexdigest()[:24]
            else:
                skey=f"AUTO-{article}-{rnum}"
        base=max(0,num(cell(row,m,"base_price",0),0))
        qty=max(0,num(cell(row,m,"quantity",1),1))
        item={
          "source_key":skey,
          "article":article,
          "nomenclature":name,
          "quantity":qty,
          "base_price":base,
          "currency":"BYN",
          "source_type":str(cell(row,m,"source_type","") or "").strip() or None,
          "source_document":doc or None,
          "source_date":iso_date(cell(row,m,"source_date"),report_date),
          "serial_no":serial or None,
          "source_comment":comment or None,
          "warranty_active":bool_warranty(cell(row,m,"warranty_active")),
          "warranty_months":int_or_none(cell(row,m,"warranty_months")),
          "warranty_note":str(cell(row,m,"warranty_note","") or "").strip() or None
        }
        rows.append(item)
    if not rows: raise RuntimeError("В файле нет строк уценки.")
    return rows

def rpc(name,payload):
    r=requests.post(
      f"{SUPABASE_URL}/rest/v1/rpc/{name}",
      headers={
        "apikey":SUPABASE_KEY,
        "Authorization":f"Bearer {SUPABASE_KEY}",
        "Content-Type":"application/json"
      },
      json=payload,timeout=120
    )
    if r.status_code>=300: raise RuntimeError(f"Supabase {r.status_code}: {r.text[:1000]}")
    return r.json() if r.text else {}

def find_latest():
    mail=imaplib.IMAP4_SSL(IMAP_HOST,IMAP_PORT)
    mail.login(IMAP_USER,IMAP_PASS)
    mail.select("INBOX")
    since=(datetime.now(MINSK)-timedelta(days=LOOKBACK_DAYS)).strftime("%d-%b-%Y")
    typ,data=mail.search(None,"SINCE",since)
    if typ!="OK": raise RuntimeError("IMAP search failed")
    best=None
    for uid in reversed(data[0].split()[-300:]):
        typ,msgdata=mail.fetch(uid,"(RFC822)")
        if typ!="OK" or not msgdata or not isinstance(msgdata[0],tuple): continue
        msg=email.message_from_bytes(msgdata[0][1])
        subject=decoded(msg.get("Subject"))
        low=subject.lower()
        if SUBJECT_KEYS and not all(k in low for k in SUBJECT_KEYS): continue
        sent=parsedate_to_datetime(msg.get("Date")) if msg.get("Date") else datetime.now(timezone.utc)
        if sent.tzinfo is None: sent=sent.replace(tzinfo=timezone.utc)
        for part in msg.walk():
            fn=decoded(part.get_filename())
            if not fn: continue
            if not fn.lower().endswith(".xlsx"): continue
            payload=part.get_payload(decode=True)
            if not payload: continue
            candidate=(sent,subject,fn,payload)
            if best is None or sent>best[0]: best=candidate
    mail.logout()
    return best

def main():
    item=find_latest()
    if not item:
        log("Нового письма «Уценка для CRM» с XLSX не найдено.")
        return 0
    sent,subject,filename,content=item
    local_sent=sent.astimezone(MINSK)
    report_date=local_sent.date()
    rows=parse_xlsx(content,report_date,filename)
    log(f"Найдено: {subject} · {filename} · {len(rows)} строк")
    out=rpc("markdown_import_snapshot_v1",{
      "p_rows":rows,
      "p_report_date":report_date.isoformat(),
      "p_source_message_at":sent.astimezone(timezone.utc).isoformat(),
      "p_subject":subject,
      "p_filename":filename,
      "p_full_snapshot":FULL_SNAPSHOT
    })
    log("✅ "+json.dumps(out,ensure_ascii=False))
    return 0

if __name__=="__main__":
    raise SystemExit(main())
