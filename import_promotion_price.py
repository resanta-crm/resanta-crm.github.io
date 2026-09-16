#!/usr/bin/env python3
# RESANTA CRM · MO2 price import from 1C email
import os, re, io, json, imaplib, email
from datetime import datetime, timezone, timedelta
from email.header import decode_header, make_header
from email.utils import parsedate_to_datetime
from zoneinfo import ZoneInfo

import requests
from openpyxl import load_workbook
import xlrd

IMAP_HOST=os.getenv("IMAP_HOST","imap.gmail.com")
IMAP_PORT=int(os.getenv("IMAP_PORT","993"))
IMAP_USER=os.environ.get("IMAP_USER","").strip()
IMAP_PASS="".join(os.environ.get("IMAP_PASS","").split()).replace("\xa0","")
SUPABASE_URL=os.environ.get("SUPABASE_URL","").strip().rstrip("/")
SUPABASE_KEY=os.environ.get("SUPABASE_KEY","").strip()
LOOKBACK_DAYS=int(os.getenv("MO2_LOOKBACK_DAYS","7"))
SUBJECT_KEYS=[x.strip().lower() for x in os.getenv("MO2_SUBJECT_KEYS","прайс мо2").split(",") if x.strip()]
MINSK=ZoneInfo("Europe/Minsk")

ALIASES={
    "sku":["артикул","код","sku","номенклатурный номер"],
    "product_name":["номенклатура","товар","наименование","наименование товара"],
    "quantity":["количество","кол-во","остаток","конечный остаток","qty"],
    "price_vat":["мелкий опт 2 с ндс","мелкий опт2 с ндс","мо2 с ндс","цена мо2 с ндс","мелкий опт 2"]
}

def log(s): print(s, flush=True)

def norm(v):
    s=str(v or "").strip().lower().replace("ё","е")
    return re.sub(r"\s+"," ",s)

def decoded(v):
    try: return str(make_header(decode_header(v or "")))
    except: return str(v or "")

def subject_matches(subject):
    s=norm(subject)
    return bool(SUBJECT_KEYS) and all(k in s for k in SUBJECT_KEYS)

def num(v, default=None):
    if v is None or v=="": return default
    if isinstance(v,(int,float)) and not isinstance(v,bool): return float(v)
    s=str(v).strip().replace("\xa0","").replace(" ","").replace(",",".")
    if not re.fullmatch(r"-?\d+(?:\.\d+)?", s): return default
    try: return float(s)
    except: return default

def choose_header_xlsx(ws):
    alias_terms={norm(x) for vals in ALIASES.values() for x in vals}
    best=None
    for r in range(1,min(ws.max_row,25)+1):
        vals=[norm(ws.cell(r,c).value) for c in range(1,min(ws.max_column,40)+1)]
        score=sum(1 for x in vals if x in alias_terms)
        if best is None or score>best[0]: best=(score,r)
    if not best or best[0]<2: raise RuntimeError("Не нашёл строку заголовков прайса МО2.")
    return best[1]

def map_headers(headers):
    m={}
    for key,aliases in ALIASES.items():
        aset={norm(x) for x in aliases}
        for idx,h in enumerate(headers):
            if norm(h) in aset:
                m[key]=idx
                break
    if "price_vat" not in m:
        candidates=[]
        for idx,h in enumerate(headers):
            nh=norm(h)
            if "мелкий опт 2" in nh and "ндс" in nh: candidates.append(idx)
        if len(candidates)==1: m["price_vat"]=candidates[0]
    for req in ("sku","product_name","quantity","price_vat"):
        if req not in m: raise RuntimeError(f"Не найдена обязательная колонка: {req}")
    return m

def dedupe(rows):
    out={}
    for row in rows:
        sku=row["sku"]
        prev=out.get(sku)
        if prev and abs(prev["price_vat"]-row["price_vat"])>0.01:
            raise RuntimeError(f"Для SKU {sku} найдены разные цены МО2: {prev['price_vat']} и {row['price_vat']}")
        out[sku]=row
    return list(out.values())

def parse_xlsx(content):
    wb=load_workbook(io.BytesIO(content),data_only=True,read_only=True)
    ws=wb[wb.sheetnames[0]]
    hr=choose_header_xlsx(ws)
    headers=[]
    for col in range(1,ws.max_column+1):
        cur=ws.cell(hr,col).value
        prev=ws.cell(hr-1,col).value if hr>1 else None
        headers.append(cur if str(cur or "").strip() else prev)
    m=map_headers(headers)
    rows=[]
    for row in ws.iter_rows(min_row=hr+1,values_only=True):
        sku=str(row[m["sku"]] or "").strip()
        name=str(row[m["product_name"]] or "").strip()
        qty=num(row[m["quantity"]],None)
        price=num(row[m["price_vat"]],None)
        if not sku or not name: continue
        # Business rule: MO2 is offered only for stock currently available in Vitebsk.
        # Rows with no stock have "Деление на 0" in 1C and become active automatically
        # on a future snapshot as soon as stock and a valid price appear.
        if qty is None or qty<=0 or price is None or price<=0: continue
        rows.append({"sku":sku,"product_name":name,"price_vat":round(price,2)})
    rows=dedupe(rows)
    if not rows: raise RuntimeError("В файле нет доступных SKU с корректной ценой МО2.")
    return rows

def parse_xls(content):
    book=xlrd.open_workbook(file_contents=content)
    sh=book.sheet_by_index(0)
    alias_terms={norm(x) for vals in ALIASES.values() for x in vals}
    best=None
    for r in range(0,min(sh.nrows,25)):
        vals=[norm(sh.cell_value(r,c)) for c in range(0,min(sh.ncols,40))]
        score=sum(1 for x in vals if x in alias_terms)
        if best is None or score>best[0]: best=(score,r)
    if not best or best[0]<2: raise RuntimeError("Не нашёл строку заголовков прайса МО2 в XLS.")
    hr=best[1]
    headers=[]
    for c in range(sh.ncols):
        cur=sh.cell_value(hr,c)
        prev=sh.cell_value(hr-1,c) if hr>0 else ""
        headers.append(cur if str(cur or "").strip() else prev)
    m=map_headers(headers)
    rows=[]
    for r in range(hr+1,sh.nrows):
        vals=[sh.cell_value(r,c) for c in range(sh.ncols)]
        sku=str(vals[m["sku"]] or "").strip()
        name=str(vals[m["product_name"]] or "").strip()
        qty=num(vals[m["quantity"]],None)
        price=num(vals[m["price_vat"]],None)
        if not sku or not name: continue
        if qty is None or qty<=0 or price is None or price<=0: continue
        rows.append({"sku":sku,"product_name":name,"price_vat":round(price,2)})
    rows=dedupe(rows)
    if not rows: raise RuntimeError("В файле нет доступных SKU с корректной ценой МО2.")
    return rows

def parse_excel(content,filename):
    return parse_xls(content) if filename.lower().endswith(".xls") else parse_xlsx(content)

def select_mailbox(mail):
    for mailbox in ['"[Gmail]/All Mail"','"[Google Mail]/All Mail"','INBOX']:
        try:
            typ,_=mail.select(mailbox,readonly=True)
            if typ=="OK":
                log(f"IMAP mailbox: {mailbox}")
                return
        except Exception: pass
    raise RuntimeError("Не удалось открыть Gmail All Mail или INBOX")

def find_latest():
    mail=imaplib.IMAP4_SSL(IMAP_HOST,IMAP_PORT)
    mail.login(IMAP_USER,IMAP_PASS)
    select_mailbox(mail)
    since=(datetime.now(MINSK)-timedelta(days=LOOKBACK_DAYS)).strftime("%d-%b-%Y")
    typ,data=mail.search(None,"SINCE",since)
    if typ!="OK":
        mail.logout(); raise RuntimeError("IMAP search failed")

    # Newest-first and return immediately on the first matching MO2 message.
    # There may be only 1-2 MO2 messages, so collecting dozens of matches would
    # waste time scanning hundreds of unrelated messages.
    ids=list(reversed(data[0].split()[-250:]))
    for uid in ids:
        typ,hdata=mail.fetch(uid,"(BODY.PEEK[HEADER.FIELDS (SUBJECT DATE FROM)])")
        if typ!="OK" or not hdata or not isinstance(hdata[0],tuple): continue
        hdr=email.message_from_bytes(hdata[0][1])
        subject=decoded(hdr.get("Subject"))
        if not subject_matches(subject): continue
        try: sent=parsedate_to_datetime(hdr.get("Date")) if hdr.get("Date") else datetime.now(timezone.utc)
        except Exception: sent=datetime.now(timezone.utc)
        if sent.tzinfo is None: sent=sent.replace(tzinfo=timezone.utc)

        typ,msgdata=mail.fetch(uid,"(RFC822)")
        if typ!="OK" or not msgdata or not isinstance(msgdata[0],tuple): continue
        msg=email.message_from_bytes(msgdata[0][1])
        for part in msg.walk():
            fn=decoded(part.get_filename())
            if not fn or not fn.lower().endswith((".xlsx",".xls")): continue
            payload=part.get_payload(decode=True)
            if payload:
                mail.logout()
                return (sent,subject,fn,payload)
    mail.logout()
    return None

def rpc(rows,report_date,sent,subject,filename):
    r=requests.post(
        f"{SUPABASE_URL}/rest/v1/rpc/promotion_price_import_snapshot_v1",
        headers={"apikey":SUPABASE_KEY,"Authorization":f"Bearer {SUPABASE_KEY}","Content-Type":"application/json"},
        json={
            "p_rows":rows,
            "p_report_date":report_date.isoformat(),
            "p_source_message_at":sent.astimezone(timezone.utc).isoformat(),
            "p_subject":subject,
            "p_filename":filename
        },timeout=120)
    if r.status_code>=300: raise RuntimeError(f"Supabase {r.status_code}: {r.text[:1500]}")
    return r.json() if r.text else {}

def main():
    if not all((IMAP_USER,IMAP_PASS,SUPABASE_URL,SUPABASE_KEY)):
        raise RuntimeError("Не заданы IMAP/Supabase secrets")
    item=find_latest()
    if not item:
        log("Нового письма «Прайс МО2» с Excel-вложением не найдено.")
        return 0
    sent,subject,filename,content=item
    report_date=sent.astimezone(MINSK).date()
    rows=parse_excel(content,filename)
    log(f"Найдено: {subject} · {filename} · {len(rows)} активных SKU")
    out=rpc(rows,report_date,sent,subject,filename)
    log("✅ "+json.dumps(out,ensure_ascii=False))
    return 0

if __name__=="__main__":
    raise SystemExit(main())
